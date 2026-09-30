"""Writes turns-deck.pptx, the corpus's synthetic deck of stored geometry
the text must leave alone (CC0-1.0), with python-pptx 1.0.2 from its default
template, the XML then set by hand the way other applications store it:

- slide 1: rotations as Google Slides and older PowerPoint files store
  them, negative (rot="-5400000", -90 degrees) or past a full turn
  (rot="25200000", 420 degrees), flips written flipH="true", a connector
  turned -90 degrees and a group whose object is turned -45 degrees;
- slide 2: a picture in mc:AlternateContent at the top of the slide (as
  PowerPoint stores a 3D model or ink with a picture for older readers),
  its box in EMU that are not whole points.

Run: pip install python-pptx==1.0.2 && python3 make_turns.py. Regenerating
gives the same content, not the same bytes (zip timestamps)."""
import copy, datetime, io
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from PIL import Image

prs = Presentation()
L = {l.name: l for l in prs.slide_layouts}


def xfrm(shape):
    return shape._element.spPr.find(qn("a:xfrm"))


# Slide 1: rotations and flips as other applications store them.
s = prs.slides.add_slide(L["Title Only"])
s.shapes.title.text = "Turned objects"
a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(60), Pt(160), Pt(200), Pt(40))
a.name = "Minus ninety"
a.text = "rot -5400000"
xfrm(a).set("rot", "-5400000")
b = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Pt(300), Pt(160), Pt(160), Pt(60))
b.name = "Minus fifteen flipped"
b.text = "rot -900000 flipH true"
xfrm(b).set("rot", "-900000")
xfrm(b).set("flipH", "true")
c = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(500), Pt(160), Pt(160), Pt(60))
c.name = "Past a turn"
c.text = "rot 25200000"
xfrm(c).set("rot", "25200000")
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(354), Pt(260), Pt(354), Pt(400))
ln.name = "Line turned"
xfrm(ln).set("rot", "-5400000")
xfrm(ln).set("flipV", "1")
g = s.shapes.add_group_shape()
g.name = "Group with a turn"
g1 = g.shapes.add_shape(MSO_SHAPE.OVAL, Pt(80), Pt(420), Pt(80), Pt(80))
g1.name = "Turned in group"
xfrm(g1).set("rot", "-2700000")
g2 = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(200), Pt(430), Pt(120), Pt(60))
g2.name = "Plain in group"

# Slide 2: a picture in mc:AlternateContent at the top of the slide.
s = prs.slides.add_slide(L["Title Only"])
s.shapes.title.text = "One object in two forms"
buf = io.BytesIO()
Image.new("RGB", (64, 64), (200, 120, 40)).save(buf, "PNG")
buf.seek(0)
pic = s.shapes.add_picture(buf, Emu(4606986), Emu(1920956), Emu(2978028), Emu(3016087))
pic.name = "Model preview"
pic._element.nvPicPr.cNvPr.set("descr", "Smiling face")
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
el = pic._element
tree = el.getparent()
at = tree.index(el)
tree.remove(el)
ac = etree.Element(etree.QName(MC, "AlternateContent"), nsmap={"mc": MC})
choice = etree.SubElement(ac, etree.QName(MC, "Choice"), nsmap={"a14": A14})
choice.set("Requires", "a14")
chosen = copy.deepcopy(el)
blip = chosen.find(".//" + qn("a:blip"))
ext = etree.SubElement(etree.SubElement(blip, qn("a:extLst")), qn("a:ext"))
ext.set("uri", "{28A0092B-C50C-407E-A947-70E740481C1C}")
etree.SubElement(ext, etree.QName(A14, "useLocalDpi")).set("val", "0")
choice.append(chosen)
etree.SubElement(ac, etree.QName(MC, "Fallback")).append(el)
tree.insert(at, ac)

cp = prs.core_properties
cp.author = "hanji"
cp.last_modified_by = "hanji"
cp.title = "Turned objects"
cp.created = cp.modified = datetime.datetime(2026, 9, 30, 0, 0, 0)
prs.save("turns-deck.pptx")
