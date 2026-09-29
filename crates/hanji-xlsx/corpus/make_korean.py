#!/usr/bin/env python3
"""Writes korean-sales.xlsx (openpyxl 3.1.5) and, when LibreOffice is installed,
korean-sales-lo.xlsx (the same workbook saved again by LibreOffice, which
computes the formulas' cached values). Both are synthetic, written for this
project (CC0-1.0). Regenerating gives the same content but not the same bytes
(zip timestamps). Run: python3 make_korean.py
"""
import os
import random
import shutil
import subprocess
import tempfile
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

random.seed(20260929)
BRANCHES = ["강남", "서초", "송파", "분당", "일산"]
LINES = ["가전", "모바일", "생활"]

wb = Workbook()
ws = wb.active
ws.title = "매출"
ws["A1"] = "2026 영업실적 (단위: 원)"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:F1")
ws["A1"].alignment = Alignment(horizontal="center")
head = ["월", "지점", "제품군", "매출", "원가", "이익"]
for c, h in enumerate(head, 1):
    ws.cell(row=3, column=c, value=h)
r = 4
for m in (1, 2, 3):
    for b in BRANCHES:
        for ln in LINES:
            sales = random.randrange(4_000, 19_000) * 1000
            cost = int(sales * random.uniform(0.66, 0.78)) // 10_000 * 10_000
            ws.cell(row=r, column=1, value=date(2026, m, 1)).number_format = "yyyy-mm"
            ws.cell(row=r, column=2, value=b)
            ws.cell(row=r, column=3, value=ln)
            ws.cell(row=r, column=4, value=sales).number_format = "#,##0"
            ws.cell(row=r, column=5, value=cost).number_format = "#,##0"
            ws.cell(row=r, column=6, value=f"=Sales[[#This Row],[매출]]-Sales[[#This Row],[원가]]").number_format = "#,##0"
            r += 1
last = r - 1
t = Table(displayName="Sales", ref=f"A3:F{last}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
for k, h in enumerate(head, 1):
    t.tableColumns.append(TableColumn(id=k, name=h))
t.tableColumns[5].calculatedColumnFormula = TableFormula(attr_text="Sales[[#This Row],[매출]]-Sales[[#This Row],[원가]]")
ws.add_table(t)
dv = DataValidation(type="list", formula1='"강남,서초,송파,분당,일산"', allow_blank=True)
dv.add(f"B4:B{last}")
ws.add_data_validation(dv)
ws.conditional_formatting.add(f"F4:F{last}", CellIsRule(operator="lessThan", formula=["3000000"], fill=PatternFill("solid", fgColor="FFC7CE")))
ws["H3"] = "출처"
ws["H4"] = "사내 포털"
ws["H4"].hyperlink = "https://example.com/sales"
ws["A3"].comment = Comment("월은 매월 1일로 적습니다.", "작성자")
chart = BarChart()
chart.title = "지점별 매출"
chart.add_data(Reference(ws, min_col=4, min_row=3, max_row=18), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, min_row=4, max_row=18))
ws.add_chart(chart, "H6")

staff = wb.create_sheet("담당자")
staff.append(["사번", "이름", "지점", "휴대전화", "입사일"])
names = ["김서연", "이민준", "박지우", "최하윤", "정도윤", "강서준", "조하은", "윤지호"]
for k, n in enumerate(names):
    staff.append([f"{random.randrange(100, 99_999):05d}", n, BRANCHES[k % 5], f"010-{random.randrange(1000, 9999)}-{random.randrange(1000, 9999)}", date(2019 + k % 6, 1 + k, 2)])
    staff.cell(row=k + 2, column=1).number_format = "@"
    staff.cell(row=k + 2, column=4).number_format = "@"
    staff.cell(row=k + 2, column=5).number_format = "yyyy-mm-dd"
t2 = Table(displayName="Staff", ref=f"A1:E{len(names) + 1}")
t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
staff.add_table(t2)

summary = wb.create_sheet("요약")
summary["A1"] = "지점"
summary["B1"] = "합계"
for k, b in enumerate(BRANCHES, 2):
    summary.cell(row=k, column=1, value=b)
    summary.cell(row=k, column=2, value=f'=SUMIFS(Sales[매출],Sales[지점],A{k})').number_format = "#,##0"
summary["A8"] = "총계"
summary["B8"] = "=SUM(B2:B6)"
summary["B8"].number_format = "#,##0"

here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, "korean-sales.xlsx")
wb.save(out)
print("wrote", out)

if shutil.which("soffice"):
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", d, out], check=True, capture_output=True)
        shutil.copy(os.path.join(d, "korean-sales.xlsx"), os.path.join(here, "korean-sales-lo.xlsx"))
        print("wrote", os.path.join(here, "korean-sales-lo.xlsx"))
